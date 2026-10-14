"""非 Markdown 文档解析器：PDF / DOCX / XLSX → Markdown 文本。

统一出口是 ``parse_to_markdown(path) -> str``，返回可直接交给
``MarkdownChunker`` 的 Markdown 文本。设计要点：

- 产物保留结构信息（PDF 按页、DOCX 按标题层级、XLSX 按 sheet 与行），
  使下游分段器能按标题切出有语义的章节。
- 对应解析库缺失时抛 ``ParserDependencyError``，由调用方决定跳过或报错，
  避免"静默丢文档"。
- 单个文件解析失败不应阻断其他文件，异常向上抛由 loader 捕获记录。
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

# 支持的非 Markdown 扩展名（小写，含点）。
SUPPORTED_EXTS = (".pdf", ".docx", ".xlsx")


class ParserDependencyError(RuntimeError):
    """对应格式的解析库未安装。"""


def is_supported(path: str | Path) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_EXTS


def parse_to_markdown(path: str | Path) -> str:
    """把 PDF / DOCX / XLSX 文件解析成 Markdown 文本。

    未知扩展名抛 ``ValueError``；依赖缺失抛 ``ParserDependencyError``；
    文件损坏等解析错误抛底层异常，由调用方捕获。
    """
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(p)
    if ext == ".docx":
        return _parse_docx(p)
    if ext == ".xlsx":
        return _parse_xlsx(p)
    raise ValueError(f"unsupported document type: {ext}")


# ---------------------------------------------------------------------- PDF
def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - 取决于安装环境
        raise ParserDependencyError("缺少 pypdf 依赖，请执行 pip install pypdf") from exc

    reader = PdfReader(str(path))
    stem = path.stem
    parts: list[str] = [f"# {stem}"]
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        parts.append(f"\n## 第 {index} 页\n")
        parts.append(text)
    if len(parts) == 1:
        # 全篇无文本（可能是扫描件图片），不产出空文档。
        logger.warning("pdf has no extractable text (scanned image?): %s", path)
    return "\n".join(parts).strip() + "\n"


# --------------------------------------------------------------------- DOCX
def _parse_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - 取决于安装环境
        raise ParserDependencyError(
            "缺少 python-docx 依赖，请执行 pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            # Word 标题 1→#、标题 2→##，映射到 Markdown 标题层级。
            level = _heading_level(para.style.name)
            parts.append("\n" + "#" * level + " " + text + "\n")
        else:
            parts.append(text)

    # 表格转成 Markdown 表格（取每行单元格文本）。
    for table in document.tables:
        rows = [_table_row_cells(row) for row in table.rows]
        rows = [cells for cells in rows if any(c for c in cells)]
        if not rows:
            continue
        width = max(len(cells) for cells in rows)
        parts.append("")
        for i, cells in enumerate(rows):
            padded = (cells + [""] * width)[:width]
            parts.append("| " + " | ".join(padded) + " |")
            if i == 0:
                parts.append("|" + "---|" * width)
        parts.append("")

    text = "\n".join(parts).strip()
    if not text:
        text = f"# {path.stem}"
    return text + "\n"


def _heading_level(style_name: str) -> int:
    digits = "".join(ch for ch in style_name if ch.isdigit())
    if digits:
        return min(6, max(1, int(digits)))
    return 2


def _table_row_cells(row) -> list[str]:
    cells = []
    for cell in row.cells:
        text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
        cells.append(text.replace("|", "\\|"))
    return cells


# --------------------------------------------------------------------- XLSX
def _parse_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - 取决于安装环境
        raise ParserDependencyError("缺少 openpyxl 依赖，请执行 pip install openpyxl") from exc

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = [f"# {path.stem}"]
    try:
        for sheet in workbook.worksheets:
            parts.append(f"\n## {sheet.title}\n")
            parts.extend(_sheet_rows_to_markdown(sheet))
    finally:
        workbook.close()
    return "\n".join(parts).strip() + "\n"


def _sheet_rows_to_markdown(sheet) -> list[str]:
    """把 sheet 的行转成 Markdown 表格；首行非空行视为表头。"""
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = [_cell_text(value) for value in row]
        if any(cells):
            rows.append(cells)
    if not rows:
        return []

    width = max(len(cells) for cells in rows)
    lines: list[str] = []
    for i, cells in enumerate(rows):
        padded = (cells + [""] * width)[:width]
        lines.append("| " + " | ".join(padded) + " |")
        if i == 0:
            lines.append("|" + "---|" * width)
    return lines


def _cell_text(value) -> str:
    if value is None:
        return ""
    # float 整数值去掉 .0，避免 "1.0" 这种噪音。
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("|", "\\|").replace("\n", " ").strip()
