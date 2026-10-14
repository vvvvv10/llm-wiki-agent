"""文档解析器测试：PDF / DOCX / XLSX → Markdown。"""

from __future__ import annotations

import pytest

from app.ingestion.parsers import (
    SUPPORTED_EXTS,
    is_supported,
    parse_to_markdown,
)


def test_is_supported():
    assert is_supported("a.pdf")
    assert is_supported("a.DOCX")
    assert is_supported("a.xlsx")
    assert not is_supported("a.md")
    assert not is_supported("a.txt")


def test_unknown_extension_raises():
    with pytest.raises(ValueError, match="unsupported"):
        parse_to_markdown("notes.txt")


# ------------------------------------------------------------------ PDF
def _minimal_pdf_with_text(lines: list[str]) -> bytes:
    """手写一个每行一个文本行的最小合法 PDF。"""
    contents = []
    y = 700
    for i, line in enumerate(lines, start=1):
        contents.append(f"BT /F1 12 Tf 72 {y} Td ({line}) Tj ET")
        y -= 30
    stream = "\n".join(contents).encode("utf-8")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_parse_pdf_scanned_image_produces_header_only(tmp_path):
    pytest.importorskip("pypdf")
    # 空白页 PDF（无文本）→ 只有标题行
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as f:
        writer.write(f)

    markdown = parse_to_markdown(path)
    assert markdown.startswith("# scan")


# ------------------------------------------------------------------ DOCX
def test_parse_docx_headings_and_tables(tmp_path):
    pytest.importorskip("docx")
    import docx

    document = docx.Document()
    document.add_heading("费用管理制度", level=1)
    document.add_heading("报销标准", level=2)
    document.add_paragraph("住宿每晚不超过 400 元。")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "职级"
    table.rows[0].cells[1].text = "上限"
    table.rows[1].cells[0].text = "员工"
    table.rows[1].cells[1].text = "400"

    path = tmp_path / "policy.docx"
    document.save(str(path))

    markdown = parse_to_markdown(path)
    assert "# 费用管理制度" in markdown
    assert "## 报销标准" in markdown
    assert "住宿每晚不超过 400 元。" in markdown
    assert "| 职级 | 上限 |" in markdown
    assert "|---|---|" in markdown
    assert "| 员工 | 400 |" in markdown


# ------------------------------------------------------------------ XLSX
def test_parse_xlsx_sheets_as_tables(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet1 = workbook.active
    sheet1.title = "报价单"
    sheet1.append(["产品", "价格"])
    sheet1.append(["灵犀标准版", 999])
    sheet2 = workbook.create_sheet("联系人")
    sheet2.append(["姓名", "电话"])
    sheet2.append(["张三", "13800000000"])

    path = tmp_path / "price.xlsx"
    workbook.save(str(path))

    markdown = parse_to_markdown(path)
    assert markdown.startswith("# price")
    assert "## 报价单" in markdown
    assert "| 产品 | 价格 |" in markdown
    assert "| 灵犀标准版 | 999 |" in markdown  # float 整数去掉 .0
    assert "## 联系人" in markdown
    assert "| 张三 | 13800000000 |" in markdown


def test_parse_xlsx_empty_sheet_skipped(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.title = "空表"
    path = tmp_path / "empty.xlsx"
    workbook.save(str(path))

    markdown = parse_to_markdown(path)
    assert "## 空表" in markdown
    assert "|" not in markdown  # 无数据行不产出表格


# ---------------------------------------------------- 加载器集成
def test_loader_picks_up_binary_documents(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl

    from app.ingestion.loader import MarkdownLoader

    root = tmp_path / "wiki"
    root.mkdir()
    (root / "guide.md").write_text("# 指南\n\nMarkdown 文档。\n", encoding="utf-8")

    workbook = openpyxl.Workbook()
    workbook.active.title = "库存"
    workbook.active.append(["物品", "数量"])
    workbook.active.append(["笔记本", 12])
    workbook.save(str(root / "stock.xlsx"))

    docs = MarkdownLoader(root).load()
    paths = sorted(d.path for d in docs)
    assert paths == ["guide.md", "stock.xlsx"]
    xlsx_doc = next(d for d in docs if d.path == "stock.xlsx")
    assert xlsx_doc.title == "stock"  # 无一级标题回退到文件名
    assert "| 笔记本 | 12 |" in xlsx_doc.content


def test_loader_survives_corrupt_binary(tmp_path):
    """损坏的二进制文档跳过，不影响其他文件加载。"""
    from app.ingestion.loader import MarkdownLoader

    root = tmp_path / "wiki"
    root.mkdir()
    (root / "ok.md").write_text("# 正常\n\n内容。\n", encoding="utf-8")
    (root / "broken.pdf").write_bytes(b"not a pdf at all")

    docs = MarkdownLoader(root).load()
    assert [d.path for d in docs] == ["ok.md"]


def test_supported_exts_constant():
    assert set(SUPPORTED_EXTS) == {".pdf", ".docx", ".xlsx"}
